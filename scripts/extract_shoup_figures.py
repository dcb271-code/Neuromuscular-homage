#!/usr/bin/env python3
"""
Extract selected tracings from Jaime Shoup, MD, "Non-Epileptiform EEG Abnormalities"
(University of Louisville lecture deck), used on the site with the author's permission
(confirmed to the site owner, 2026-09-30).

Slides were reviewed one by one. Excluded: slides carrying a Mayo Foundation permission
line (35, 37, 39), a scanned journal figure (24), an unlabelled tracing (5), a stock
photograph (42) and a GIF of uncertain origin (3).

    .venv/bin/python scripts/extract_shoup_figures.py   # needs docs/eeg/sources/Nonepileptiform Abn.pptx

Writes public/eeg/figures/shoup-sNN.jpg and merges entries into src/eeg/figures.json.
Run scripts/build_eeg_attributions.py afterwards.
"""
from __future__ import annotations

import io
import json
from pathlib import Path

from PIL import Image
from pptx import Presentation

SRC = Path("docs/eeg/sources/Nonepileptiform Abn.pptx")
OUT_IMG = Path("public/eeg/figures")
REG = Path("src/eeg/figures.json")
CREDIT = "Jaime Shoup, MD, University of Louisville; used with permission"
LICENCE = "Used with permission"
MAX_W = 1600

# slide number → (title, caption). Captions describe what the tracing shows and the
# lecturer's annotations; they add no numbers beyond those written on the image.
SLIDES = {
    1: ("Left hemispheric polymorphic delta slowing",
        "Continuous, irregular delta activity over the left hemisphere (boxed), with a better organized right hemisphere. ECG channel at the bottom."),
    6: ("Focal slowing of the left hemisphere",
        "The left side is slower and less organized, with polymorphic theta to delta activity and a poorly sustained posterior rhythm; the right side is well organized, with alpha activity over the posterior regions (lecturer's annotations)."),
    10: ("Mild generalized slowing",
         "Slow, poorly sustained posterior dominant rhythm, poor anterior-to-posterior gradient, and excess theta admixed with the alpha (lecturer's annotations)."),
    11: ("Moderate generalized slowing",
         "Very slow posterior dominant rhythm, no anterior-to-posterior gradient, minimal variability, and diffuse theta and delta activity (lecturer's annotations)."),
    12: ("Severe generalized slowing",
         "Diffuse delta with no posterior dominant rhythm, no anterior-to-posterior gradient and no reactivity (lecturer's annotation)."),
    15: ("Frontal intermittent rhythmic delta activity (FIRDA)",
         "Bursts of rhythmic delta with a bifrontal maximum (arrows)."),
    18: ("Occipital intermittent rhythmic delta activity (OIRDA)",
         "Runs of rhythmic delta over both occipital regions (boxed), a childhood pattern."),
    20: ("Temporal intermittent rhythmic delta activity (TIRDA)",
         "Left temporal rhythmic delta, most rhythmic anteriorly and becoming polymorphic as it moves posteriorly (lecturer's annotation)."),
    26: ("Generalized attenuation",
         "Very low amplitude activity in every channel at a 140 µV scale; the ECG channel is the only conspicuous signal."),
    28: ("Diffuse beta activity with ECG artifact",
         "Transverse montage. Generalized fast activity of the kind seen with benzodiazepines, and a time-locked ECG artifact (lecturer's labels)."),
    30: ("Burst suppression",
         "Longitudinal bipolar montage. Bursts of mixed-frequency activity separated by suppressed intervals."),
    33: ("Alpha coma pattern (lecture example)",
         "Referential montage with widespread, monotonous alpha-range activity, shown in the lecture as an alpha coma pattern in a comatose patient."),
}


def main() -> None:
    prs = Presentation(SRC)
    OUT_IMG.mkdir(parents=True, exist_ok=True)
    for old in OUT_IMG.glob("shoup-s*.jpg"):
        old.unlink()
    entries = []
    for n, slide in enumerate(prs.slides, 1):
        if n not in SLIDES:
            continue
        pics = [s for s in slide.shapes if getattr(s, "image", None) is not None]
        if not pics:
            raise SystemExit(f"slide {n}: no image")
        img = Image.open(io.BytesIO(pics[0].image.blob)).convert("RGB")
        if img.width > MAX_W:
            img = img.resize((MAX_W, round(img.height * MAX_W / img.width)), Image.LANCZOS)
        fid = f"shoup-s{n:02d}"
        img.save(OUT_IMG / f"{fid}.jpg", "JPEG", quality=85, optimize=True)
        title, caption = SLIDES[n]
        entries.append({
            "id": fid, "number": n, "title": title, "caption": f"{title}. {caption}", "page": n,
            "file": f"/eeg/figures/{fid}.jpg", "width": img.width, "height": img.height,
            "credit": CREDIT, "licence": LICENCE, "source": "shoup",
        })
    reg = [r for r in json.loads(REG.read_text()) if not r["id"].startswith("shoup-")]
    for r in reg:
        r.setdefault("source", "aes")
    reg.extend(entries)
    REG.write_text(json.dumps(reg, indent=1, ensure_ascii=False))
    print(f"wrote {len(entries)} Shoup figures; registry now {len(reg)} entries")


if __name__ == "__main__":
    main()
